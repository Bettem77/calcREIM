# -*- coding: utf-8 -*-
"""
calc_eta_star.py
DSA(動的信号解析器)の生データCSV([Header]/[Calibration]/[TIME_INST] 形式)を、
eta_star_v5 形式の計算用xlsxの「TimeData」シートに自動転記する。
(手作業での「A列=時刻, B列=力, D列=加速度 を貼り付け」を代わりに行うスクリプト)

入力 : RAWDATA_DIR (既定 "Rawdata" フォルダ) 内の <label>_<f>Hz_<level>_<date>.csv
        例) Rawdata/100um_50Hz_5_0926.csv
        - [Calibration] の "EU/V" 行から、各チャンネルの物理量(mm/g/m/s2/N など)への
          換算係数[EU/V]を読み取る(電圧 × 係数 = 物理量)。
        - [TIME_INST] の CH電圧のうち、単位が "N" のチャンネルを力、"m/s2" のチャンネルを
          加速度として使う(チャンネル番号を直接は決め打ちしない)。

出力 : ETASTAR_DIR (既定 "eta_star" フォルダ) 内に eta_star_v5_<label>_<f>Hz_<level>_<date>.xlsx
        (生データCSVの日付をそのまま末尾に付けるので、同じ label/f/level の過去データと
         ファイル名が衝突しない。effective_mass_batch.py / force_waveform_batch.py 側の
         ファイル名パターンも、この末尾日付を許容するよう対応済み)
        TEMPLATE_XLSX(eta_star フォルダ内の eta_star_v5_box_1000Hz_10.xlsx)を複製し、
        TimeData!A/B/D と EtaStar_Acc!B6(f) を書き換える。
        既定では出力先が既に存在する場合は上書きしない(--overwrite を付けると上書きする)。

        ※ 粉体質量 m_p [EtaStar_Acc!B7] は生データから分からないため、実行時に粒径(label)ごとに
          対話入力で確認する(同じ粒径のファイルはまとめて1回だけ質問、空欄でスキップも可)。
          容器質量 m_c [EtaStar_Acc!B8] は既定値 M_C_DEFAULT を仮入力する。
        ※ 使用するデータ点数 N も実行時に周波数(f)ごとに対話入力できる
          (先頭から何点使うか。空欄なら全点使用)。
"""
import re
import glob
import os
import sys
from openpyxl import load_workbook

# ===================== 設定 =====================
INPUT_DIR = r"C:\Users\USER\Desktop\calcREIM"
RAWDATA_DIR = os.path.join(INPUT_DIR, "Rawdata")    # 生データCSVはここに置く
ETASTAR_DIR = os.path.join(INPUT_DIR, "eta_star")   # eta_star_v5_*.xlsx の出力/参照先
TEMPLATE_XLSX = os.path.join(ETASTAR_DIR, "eta_star_v5_box_1000Hz_10.xlsx")
M_C_DEFAULT = 0.88          # 容器(ジグ)質量[g]。プロジェクト内でほぼ一貫している既定値。要確認。
FORCE_SIGN = -1             # 力チャンネルの符号補正。センサー配線の都合で "EU/V × 電圧" の
                            # そのままの符号だと、既存の正しいxlsx(F と a が正相関)と逆に
                            # なる(実測で相関 +0.97 → -0.98 と符号反転を確認済み)。
FORCE_UNIT = "N"            # このラベルを持つチャンネルを力として使う
ACC_UNIT_RE = re.compile(r"m/s\^?2", re.IGNORECASE)  # "m/s2" "m/s^2" どちらも許容
# ===============================================

RAW_RE = re.compile(
    r"^(?P<label>.+)_(?P<f>\d+(?:\.\d+)?)Hz_(?P<level>[^_]+)_(?P<date>\d+)\.csv$",
    re.IGNORECASE)
CH_RE = re.compile(r"CH(\d+)\((.+)\)")
CH_V_RE = re.compile(r"CH(\d+)\(V\)", re.IGNORECASE)


def parse_calibration(lines):
    """[Calibration] セクションから チャンネル番号 -> (単位, 係数[EU/V]) を返す"""
    for i, line in enumerate(lines):
        if line.strip() == "[Calibration]":
            header = None
            for j in range(i + 1, len(lines)):
                cells = [c.strip() for c in lines[j].rstrip("\r\n").split(",")]
                if not cells or cells[0] != "EU/V":
                    if header is not None:
                        break
                    continue
                if header is None:
                    header = cells  # 例: ['EU/V','CH1(mm)','CH2(g)','CH3(m/s2)','CH4(N)','']
                    continue
                factors = cells
                ch_info = {}
                for k in range(1, len(header)):
                    m = CH_RE.match(header[k])
                    if m and k < len(factors) and factors[k]:
                        ch_info[int(m.group(1))] = (m.group(2), float(factors[k]))
                return ch_info
            break
    raise ValueError("[Calibration] の EU/V 行が見つからない")


def find_channel(ch_info, unit_matcher):
    for ch, (unit, factor) in ch_info.items():
        if (unit_matcher == FORCE_UNIT and unit == FORCE_UNIT) or \
           (unit_matcher == "ACC" and ACC_UNIT_RE.fullmatch(unit)):
            return ch, factor
    raise ValueError(f"単位 {unit_matcher!r} のチャンネルが見つからない: {ch_info}")


def read_time_inst(lines, ch_time_col, ch_force, factor_force, ch_acc, factor_acc):
    t, F, a = [], [], []
    in_block = False
    col_of_ch = None
    for line in lines:
        s = line.rstrip("\r\n")
        if s.strip() == "[TIME_INST]":
            in_block = True
            continue
        if not in_block:
            continue
        cells = [c.strip() for c in s.split(",")]
        if cells[0].startswith("Time(s)"):
            col_of_ch = {}
            for k, c in enumerate(cells):
                m = CH_V_RE.match(c)
                if m:
                    col_of_ch[int(m.group(1))] = k
            continue
        if col_of_ch is None:
            continue
        try:
            ti = float(cells[0])
        except ValueError:
            break  # "[EOF]" や空行に到達
        Fi = float(cells[col_of_ch[ch_force]]) * factor_force * FORCE_SIGN
        ai = float(cells[col_of_ch[ch_acc]]) * factor_acc
        t.append(ti); F.append(Fi); a.append(ai)
    if not t:
        raise ValueError("[TIME_INST] のデータが読めなかった")
    return t, F, a


def read_raw_csv(path):
    with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
        lines = f.readlines()
    ch_info = parse_calibration(lines)
    ch_force, factor_force = find_channel(ch_info, FORCE_UNIT)
    ch_acc, factor_acc = find_channel(ch_info, "ACC")
    t, F, a = read_time_inst(lines, None, ch_force, factor_force, ch_acc, factor_acc)
    return t, F, a


def ask_m_p(label):
    """粒径labelごとに粉体質量 m_p [g] を対話入力で聞く。空欄ならNone(未入力)のまま返す。"""
    while True:
        raw = input(f"  [{label}] 粉体質量 m_p [g] を入力(空欄でスキップ): ").strip()
        if raw == "":
            return None
        try:
            return float(raw)
        except ValueError:
            print("    数値を入力してください(例: 0.804)。空欄で未入力のまま進めます。")


def ask_n_points(f_str):
    """周波数f_strごとに 使用するデータ点数 N を対話入力で聞く。空欄なら全点使用(None)を返す。"""
    while True:
        raw = input(f"  [{f_str}Hz] 使用するデータ点数 N (先頭から何点使うか。空欄で全点使用): ").strip()
        if raw == "":
            return None
        try:
            n = int(raw)
            if n <= 0:
                raise ValueError
            return n
        except ValueError:
            print("    正の整数を入力してください。空欄なら全点使用します。")


def write_xlsx(label, f_str, level, t, F, a, out_path, m_p):
    wb = load_workbook(TEMPLATE_XLSX)  # 数式を保持するため data_only=False (既定)
    ws = wb["TimeData"]

    # 旧データの取り残し防止: 新データより長い範囲までクリア
    clear_to = max(ws.max_row, len(t) + 1) + 5
    for row in ws.iter_rows(min_row=2, max_row=clear_to, min_col=1, max_col=4):
        for c in row:
            c.value = None

    for i, (ti, Fi, ai) in enumerate(zip(t, F, a), start=2):
        ws.cell(row=i, column=1, value=ti)   # A: 時刻[s]
        ws.cell(row=i, column=2, value=Fi)   # B: 力[N]
        ws.cell(row=i, column=4, value=ai)   # D: 加速度[m/s^2]

    ws_acc = wb["EtaStar_Acc"]
    ws_acc["B6"] = float(f_str)   # 加振周波数(ファイル名から自動設定)
    ws_acc["B7"] = m_p            # 粉体質量 m_p: 実行時に対話入力(未入力ならNoneのまま空欄)
    ws_acc["B8"] = M_C_DEFAULT    # 容器質量 m_c: 既定値(要確認)

    wb.calculation.fullCalcOnLoad = True  # Excelで開いたときに強制再計算
    wb.save(out_path)


def main():
    overwrite_ok = "--overwrite" in sys.argv
    files = sorted(glob.glob(os.path.join(RAWDATA_DIR, "*.csv")))
    targets = []
    for path in files:
        m = RAW_RE.match(os.path.basename(path))
        if not m:
            continue
        targets.append((path, m))
    if not targets:
        print("対象の生データCSVが見つからない(命名規則: <label>_<f>Hz_<level>_<date>.csv)")
        return

    print(f"テンプレート: {TEMPLATE_XLSX}")

    # 上書きしない分は先に除外してから、実際に処理する粒径だけ質問する
    to_process = []
    skipped = []
    for path, m in targets:
        label, f_str, level, date = m.group("label"), m.group("f"), m.group("level"), m.group("date")
        out_name = f"eta_star_v5_{label}_{f_str}Hz_{level}_{date}.xlsx"
        out_path = os.path.join(ETASTAR_DIR, out_name)
        existed = os.path.exists(out_path)
        if existed and not overwrite_ok:
            print(f"[スキップ] {out_name} は既に存在するため上書きしない"
                  "(上書きするには --overwrite を付けて実行)")
            skipped.append(out_name)
            continue
        to_process.append((path, label, f_str, level, out_name, out_path, existed))

    if not to_process:
        return

    labels = sorted({label for _, label, *_ in to_process})
    print("\n粉体質量 m_p [g] を粒径ごとに入力してください(同じ粒径のファイルには同じ値を使う):")
    m_p_by_label = {label: ask_m_p(label) for label in labels}

    freqs = sorted({f_str for _, _, f_str, *_ in to_process}, key=lambda s: float(s))
    print("\n使用するデータ点数 N を周波数ごとに入力してください(同じ周波数のファイルには同じ値を使う):")
    n_points_by_f = {f_str: ask_n_points(f_str) for f_str in freqs}

    need_mp = []
    print()
    for path, label, f_str, level, out_name, out_path, existed in to_process:
        m_p = m_p_by_label[label]
        n_points = n_points_by_f[f_str]
        t, F, a = read_raw_csv(path)
        if n_points is not None:
            t, F, a = t[:n_points], F[:n_points], a[:n_points]
        write_xlsx(label, f_str, level, t, F, a, out_path, m_p)

        tag = "[上書き]" if existed else "[新規]"
        print(f"{tag} {os.path.basename(path)}  ->  {out_name}  (N={len(t)}点, m_p={m_p})")
        if m_p is None:
            need_mp.append(out_name)

    if need_mp:
        print("\n※ 以下のファイルは m_p が未入力のままです。Excelで開いて"
              "EtaStar_Acc!B7 の黄色セルに入力してください"
              f"(m_c は既定値 {M_C_DEFAULT} を仮入力済み、必要なら修正):")
        for name in need_mp:
            print(f"  - {name}")


if __name__ == "__main__":
    main()
