# -*- coding: utf-8 -*-
"""
force_waveform_batch.py
eta_star_v5 形式の xlsx の TimeData (A=時刻[s], B=力[N]) から、力波形をそのままグラフ化する。
計算・差し引き・規格化は一切しない。

出力 : OUTPUT_DIR/waveform_<label>.png  (粒径ごと。行=周波数, 列=加速度レベル)
"""
import re, glob, os
from openpyxl import load_workbook
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

# ===================== 設定 =====================
INPUT_DIR = r"C:\Users\USER\Desktop\calcREIM"
OUTPUT_DIR = "./output_waveform"
N_CYCLES = 3          # 先頭から何周期分を描くか (時間軸 = N_CYCLES / f)
# ===============================================

FNAME_RE = re.compile(
    r"eta_star_v\d+_(?P<label>.+)_(?P<f>\d+(?:\.\d+)?)Hz_(?P<level>[^_.]+)\.xlsx$", re.IGNORECASE)


def read_one(path):
    m = FNAME_RE.search(os.path.basename(path))
    if not m:
        return None
    wb = load_workbook(path, data_only=True, read_only=True)
    t, F = [], []
    for row in wb["TimeData"].iter_rows(min_row=2, max_col=2, values_only=True):
        if isinstance(row[0], (int, float)) and isinstance(row[1], (int, float)):
            t.append(row[0]); F.append(row[1])
    wb.close()
    return dict(label=m.group("label"), f=float(m.group("f")), level=m.group("level"), t=t, F=F)


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    rs = [r for r in (read_one(p) for p in sorted(glob.glob(os.path.join(INPUT_DIR, "*.xlsx")))) if r]
    if not rs:
        print("対象ファイルなし"); return
    for label in sorted({r["label"] for r in rs}):
        sub = [r for r in rs if r["label"] == label]
        fs = sorted({r["f"] for r in sub})
        lvs = sorted({r["level"] for r in sub}, key=lambda x: float(x) if x.replace(".", "").isdigit() else x)
        fig, ax = plt.subplots(len(fs), len(lvs), figsize=(max(3.4 * len(lvs), 9), 1.9 * len(fs) + 0.8),
                               squeeze=False)
        for r in sub:
            i, j = fs.index(r["f"]), lvs.index(r["level"])
            t_end = r["t"][0] + N_CYCLES / r["f"]
            tt = [(x - r["t"][0]) * 1000 for x in r["t"] if x <= t_end]
            ax[i, j].plot(tt, r["F"][:len(tt)], c="k", lw=0.9)
            ax[i, j].grid(alpha=.3); ax[i, j].tick_params(labelsize=6)
        for i, f in enumerate(fs):
            ax[i, 0].set_ylabel(f"{f:g} Hz\nF [N]", fontsize=8)
        for j, lv in enumerate(lvs):
            ax[0, j].set_title(f"level {lv}", fontsize=9)
            ax[-1, j].set_xlabel("t [ms]")
        fig.suptitle(f"{label}: measured force F (TimeData column B), first {N_CYCLES} cycles", fontsize=10)
        fig.tight_layout()
        out = os.path.join(OUTPUT_DIR, f"waveform_{label}.png")
        fig.savefig(out, dpi=120); plt.close(fig)
        print(f"保存: {out}")


if __name__ == "__main__":
    main()
